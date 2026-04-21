"""
seed_qdrant.py

Reads data/Inventory.xlsx Inventory sheet, embeds each row
via Ollama API embeddings with bge-m3, and upserts into Qdrant collection
inventory_products with cosine distance, dimension 1024.

Idempotent by product_id hash. Uses a seed cache file that is gitignored.
"""

import hashlib
import json
import os
import sys

import openpyxl
import requests
from qdrant_client import QdrantClient
from qdrant_client.http.models import (
    Distance,
    PointStruct,
    VectorParams,
)

DATA_PATH = os.path.join(
    os.path.dirname(__file__), "..", "data", "Inventory.xlsx"
)
CACHE_PATH = os.path.join(os.path.dirname(__file__), "..", ".seed_cache.json")
COLLECTION_NAME = "inventory_products"
VECTOR_DIM = 1024

OLLAMA_HOST = os.environ.get("OLLAMA_HOST", "http://localhost:11434")
QDRANT_HOST = os.environ.get("QDRANT_HOST", "http://localhost:6333")


def load_inventory() -> list[dict]:
    wb = openpyxl.load_workbook(DATA_PATH, read_only=True)
    ws = wb["Inventory"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, row)))
    wb.close()
    return rows


def load_cache() -> dict:
    if os.path.exists(CACHE_PATH):
        with open(CACHE_PATH, "r") as f:
            return json.load(f)
    return {}


def save_cache(cache: dict):
    with open(CACHE_PATH, "w") as f:
        json.dump(cache, f, indent=2)


def row_hash(row: dict) -> str:
    """Deterministic hash of a row for change detection."""
    raw = json.dumps(row, sort_keys=True, default=str)
    return hashlib.sha256(raw.encode()).hexdigest()


def product_id_to_point_id(product_id: str) -> int:
    """Deterministic point ID from product_id."""
    h = hashlib.md5(product_id.encode()).hexdigest()
    return int(h[:15], 16)


def get_embedding(text: str) -> list[float]:
    """Get embedding from Ollama bge-m3 model."""
    resp = requests.post(
        f"{OLLAMA_HOST}/api/embeddings",
        json={"model": "bge-m3", "prompt": text},
        timeout=60,
    )
    resp.raise_for_status()
    return resp.json()["embedding"]


def ensure_collection(client: QdrantClient):
    """Create collection if it does not exist."""
    collections = [c.name for c in client.get_collections().collections]
    if COLLECTION_NAME not in collections:
        client.create_collection(
            collection_name=COLLECTION_NAME,
            vectors_config=VectorParams(size=VECTOR_DIM, distance=Distance.COSINE),
        )
        print(f"Created collection '{COLLECTION_NAME}'")
    else:
        print(f"Collection '{COLLECTION_NAME}' already exists")


def main():
    inventory = load_inventory()
    cache = load_cache()
    client = QdrantClient(url=QDRANT_HOST, timeout=30)

    ensure_collection(client)

    skipped = 0
    upserted = 0

    for row in inventory:
        pid = row.get("Product_ID", "")
        rh = row_hash(row)

        if cache.get(pid) == rh:
            skipped += 1
            continue

        # Build embedding text from product name and category
        embed_text = f"{row.get('Product_Name', '')} {row.get('Category', '')}"
        embedding = get_embedding(embed_text)

        point_id = product_id_to_point_id(pid)
        payload = {
            "product_id": pid,
            "product_name": row.get("Product_Name"),
            "category": row.get("Category"),
            "unit_price": row.get("Unit_Price"),
            "discount_pct": row.get("Discount_Pct"),
            "stock_quantity": row.get("Stock_Quantity"),
            "stock_status": row.get("Stock_Status"),
            "next_arrival_date": str(row.get("Next_Arrival_Date")) if row.get("Next_Arrival_Date") else None,
        }

        client.upsert(
            collection_name=COLLECTION_NAME,
            points=[PointStruct(id=point_id, vector=embedding, payload=payload)],
        )

        cache[pid] = rh
        upserted += 1
        print(f"Upserted {pid}: {row.get('Product_Name')}")

    save_cache(cache)
    print(f"\nDone. Upserted: {upserted}, Skipped (unchanged): {skipped}")


if __name__ == "__main__":
    main()
