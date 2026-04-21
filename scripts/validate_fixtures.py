"""
validate_fixtures.py

Loads fixtures/emails.jsonl and fixtures/expected_outputs.jsonl,
runs matcher_reference against each, and asserts the expected route matches.
"""

import json
import os
import sys

# Add scripts dir to path
sys.path.insert(0, os.path.join(os.path.dirname(__file__)))

from matcher_reference import run_matcher

# Load inventory and deliveries from Excel
import openpyxl

INV_PATH = os.path.join(
    os.path.dirname(__file__), "..", "data", "Inventory.xlsx"
)
DEL_PATH = os.path.join(
    os.path.dirname(__file__), "..", "data", "Deliveries.xlsx"
)
EMAILS_PATH = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "emails.jsonl"
)
EXPECTED_PATH = os.path.join(
    os.path.dirname(__file__), "..", "fixtures", "expected_outputs.jsonl"
)


def load_excel_data():
    inv_wb = openpyxl.load_workbook(INV_PATH, read_only=True)
    inv_ws = inv_wb["Inventory"]
    inv_headers = [cell.value for cell in next(inv_ws.iter_rows(min_row=1, max_row=1))]
    inventory = []
    for row in inv_ws.iter_rows(min_row=2, values_only=True):
        inventory.append(dict(zip(inv_headers, row)))
    inv_wb.close()

    del_wb = openpyxl.load_workbook(DEL_PATH, read_only=True)
    del_ws = del_wb["Active_Deliveries"]
    del_headers = [cell.value for cell in next(del_ws.iter_rows(min_row=1, max_row=1))]
    deliveries = []
    for row in del_ws.iter_rows(min_row=2, values_only=True):
        deliveries.append(dict(zip(del_headers, row)))
    del_wb.close()

    return inventory, deliveries


def main():
    inventory, deliveries = load_excel_data()

    with open(EMAILS_PATH, "r") as f:
        emails = [json.loads(line) for line in f if line.strip()]

    with open(EXPECTED_PATH, "r") as f:
        expected_outputs = [json.loads(line) for line in f if line.strip()]

    assert len(emails) == len(expected_outputs), (
        f"Fixture count mismatch: {len(emails)} emails vs {len(expected_outputs)} expected outputs"
    )

    passed = 0
    failed = 0

    for i, (email, expected) in enumerate(zip(emails, expected_outputs)):
        result = run_matcher(
            intent=email.get("intent", "unknown"),
            product_name=email.get("product_name"),
            customer_name=email.get("customer_name"),
            inventory=inventory,
            deliveries=deliveries,
        )

        expected_route = expected.get("route")
        actual_route = result.get("route")

        if actual_route == expected_route:
            passed += 1
        else:
            failed += 1
            print(
                f"FAIL fixture {i + 1}: expected route '{expected_route}', "
                f"got '{actual_route}' for product '{email.get('product_name')}'"
            )

        # Also check match_tier if specified
        expected_tier = expected.get("match_tier")
        if expected_tier and result.get("match_tier") != expected_tier:
            print(
                f"  WARN fixture {i + 1}: expected tier '{expected_tier}', "
                f"got '{result.get('match_tier')}'"
            )

    print(f"\nResults: {passed} passed, {failed} failed, {passed + failed} total")
    if failed > 0:
        sys.exit(1)
    else:
        print("All fixture validations passed.")


if __name__ == "__main__":
    main()
